"""Thin local routing, page-plan validation, and final PPTX QA. No LLM/API/server."""
from __future__ import annotations
import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
RENDERERS = {
    'ppt-master': 'ppt-master/skills/ppt-master/SKILL.md',
    'dashi': 'dashi-ppt-skill/skills/dashi-ppt/SKILL.md',
    'academic': 'cn-academic-spark/skills/CN_Spark_paper2ppt/SKILL.md',
    'presentation-skill': 'presentation-skill/SKILL.md',
}

def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))

def write(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def safe_project(name):
    if not re.fullmatch(r'[\w\-\u4e00-\u9fff]+', name):
        raise ValueError('Use a project name containing letters, digits, Chinese, dash or underscore')
    return ROOT / 'projects' / name

def route(args):
    text = args.prompt.lower()
    refs = [str(Path(p).resolve()) for p in args.reference]
    for p in refs:
        if not Path(p).is_file() or Path(p).suffix.lower() != '.pptx':
            raise ValueError(f'Reference must be a real local PPTX: {p}')
    kind = args.kind
    if kind == 'auto':
        if re.search(r'benchmark|latency|throughput|speedup|ablation|性能数据|实验结果|基准测试', text):
            kind = 'data'
        elif re.search(r'论文|组会|答辩|开题|文献综述|研究计划|科研报告', text):
            kind = 'academic'
        elif re.search(r'课程|作业|课堂', text) and not re.search(r'快速|快点|快一点|明天', text):
            kind = 'simple-course'
        elif re.search(r'快速|快点|快一点|明天|调研|行业研究|商业分析|竞品|内部报告|项目复盘', text):
            kind = 'business'
        else:
            kind = 'formal'
    requested = args.renderer or ('ppt-master' if refs else {
        'data':'presentation-skill', 'academic':'academic', 'business':'dashi',
        'simple-course':'ppt-master', 'formal':'ppt-master'}[kind])
    primary, reason = requested, '用户明确选择' if args.renderer else '统一路由规则'
    # Approval of a configured provider is an explicit session fact, never guessed
    # from mere presence of an API key. We neither scan credentials nor call APIs.
    if requested == 'academic' and not args.academic_ready:
        primary = 'ppt-master'
        reason = 'Academic上游强制TechnicalRoute A/B；尚未确认其图像provider和必需工作流可用，整套回退ppt-master'
    selected = ROOT / 'skills-src' / RENDERERS[primary]
    if not selected.is_file():
        if primary == 'ppt-master': raise FileNotFoundError(selected)
        primary, reason = 'ppt-master', f'{primary}未安装，整套回退ppt-master'
        selected = ROOT / 'skills-src' / RENDERERS[primary]
    result = dict(prompt=args.prompt,kind=kind,requested_renderer=requested,
        primary_renderer=primary,preprocessor='deck-dna' if refs and primary=='ppt-master' else None,
        references=refs,reason=reason,skill=str(selected),target_slide_count=args.slides,
        academic_conditions_verified=bool(args.academic_ready))
    if args.project:
        project = safe_project(args.project)
        project.mkdir(parents=True, exist_ok=True)
        (ROOT/'output'/args.project).mkdir(parents=True, exist_ok=True)
        write(project/'route.json', result)
    return result

def validate_plan(path):
    plan = read(path)
    top = ['deck_goal','audience','duration','target_slide_count','language','narrative_arc','slides']
    fields = ['slide_number','role','title','core_message','required_evidence','required_source_material','visual_strategy','must_keep','can_drop','notes']
    problems = [f'Missing {k}' for k in top if k not in plan]
    slides = plan.get('slides', [])
    if not isinstance(slides, list) or not slides:
        problems.append('slides must be a nonempty list')
        slides = []
    if plan.get('target_slide_count') != len(slides): problems.append('Slide count differs from plan')
    for i,s in enumerate(slides,1):
        for k in fields:
            if k not in s: problems.append(f'Slide {i}: missing {k}')
        if s.get('slide_number') != i: problems.append(f'Slide {i}: order mismatch')
        for k in ['role','title','core_message','visual_strategy']:
            if not isinstance(s.get(k),str) or not s[k].strip(): problems.append(f'Slide {i}: empty {k}')
        for k in ['required_evidence','required_source_material','must_keep','can_drop']:
            if not isinstance(s.get(k),list): problems.append(f'Slide {i}: {k} must be array')
    return {'status':'failed' if problems else 'passed','slide_count':len(slides),'errors':problems}

def package_qa(path, expected):
    from pptx import Presentation
    from PIL import Image
    import io
    errors, warnings = [], []
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad: errors.append(f'Bad ZIP member: {bad}')
        for required in ['[Content_Types].xml','ppt/presentation.xml','_rels/.rels']:
            if required not in z.namelist(): errors.append(f'Missing package part: {required}')
        for name in z.namelist():
            if name.endswith(('.xml','.rels')): ET.fromstring(z.read(name))
        if any(n.lower().endswith('vbaproject.bin') for n in z.namelist()):
            errors.append('Unexpected macro project')
    prs = Presentation(path)
    w,h = prs.slide_width, prs.slide_height
    if expected and len(prs.slides) != expected: errors.append(f'Expected {expected} slides, found {len(prs.slides)}')
    if abs(w/h-16/9) > .015: warnings.append({'type':'aspect_ratio','value':round(w/h,4)})
    placeholder = re.compile(r'lorem ipsum|click to (?:edit|add)|点击(?:此处)?(?:添加|编辑)|待填写|占位符|\bTODO\b|\bTBD\b|\{\{[^}]+\}\}',re.I)
    stats = []
    def flatten(shapes):
        for shape in shapes:
            if shape.shape_type == 6:
                yield from flatten(shape.shapes)
            else: yield shape
    for i,slide in enumerate(prs.slides,1):
        shapes = list(flatten(slide.shapes))
        texts = []
        for shape in shapes:
            text = shape.text if shape.has_text_frame else ''
            if shape.has_table:
                text += '\n'.join(c.text for row in shape.table.rows for c in row.cells)
            if text:
                texts.append((shape,text))
                if placeholder.search(text): errors.append(f'Slide {i}: placeholder {text[:100]}')
                if '\ufffd' in text: errors.append(f'Slide {i}: replacement character')
            # intentional bleed of full-slide backgrounds is not a content error
            if shape.width < w*.97 and (shape.left < -45720 or shape.top < -45720 or shape.left+shape.width>w+45720 or shape.top+shape.height>h+45720):
                warnings.append({'slide':i,'type':'out_of_bounds','shape':shape.name})
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.strip() and r.font.size and r.font.size.pt < 9:
                            warnings.append({'slide':i,'type':'small_text','pt':r.font.size.pt,'text':r.text[:60]})
            if shape.shape_type == 13 and not any([shape.crop_left,shape.crop_right,shape.crop_top,shape.crop_bottom]):
                try:
                    with Image.open(io.BytesIO(shape.image.blob)) as im:
                        ratio = (shape.width/shape.height)/(im.width/im.height)
                    if abs(ratio-1)>.12: warnings.append({'slide':i,'type':'image_aspect','shape':shape.name,'ratio':round(ratio,3)})
                except Exception: pass  # EMF/SVG not readable by Pillow: inspect rendered slide
        for a,(s,t) in enumerate(texts):
            for q,u in texts[a+1:]:
                ix=max(0,min(s.left+s.width,q.left+q.width)-max(s.left,q.left))
                iy=max(0,min(s.top+s.height,q.top+q.height)-max(s.top,q.top))
                small=min(s.width*s.height,q.width*q.height)
                if small and ix*iy/small>.2:
                    warnings.append({'slide':i,'type':'text_overlap_candidate','shapes':[s.name,q.name]})
        stats.append({'slide':i,'text_objects':len(texts),'characters':sum(len(t) for _,t in texts),'objects':len(shapes)})
    return dict(status='failed' if errors else 'passed',errors=errors,warnings=warnings,
        slide_count=len(prs.slides),aspect_ratio=round(w/h,5),slides=stats)

def contact_sheet(images, path):
    from PIL import Image, ImageOps, ImageDraw
    columns = min(3,len(images))
    cellw, cellh = 480,300
    sheet = Image.new('RGB',(columns*cellw, ((len(images)+columns-1)//columns)*cellh),'#e7e5e1')
    draw=ImageDraw.Draw(sheet)
    for i,p in enumerate(images):
        with Image.open(p) as im:
            thumb=ImageOps.contain(im.convert('RGB'),(456,260))
            x=(i%columns)*cellw+12; y=(i//columns)*cellh+25
            sheet.paste(thumb,(x,y))
            draw.text((x,y-17),str(i+1),fill='#202020')
    sheet.save(path,quality=91)

def qa(args):
    path=Path(args.pptx).resolve(); out=Path(args.out).resolve()
    out.mkdir(parents=True,exist_ok=True)
    result=package_qa(path,args.expected_slides)
    result.update(source=str(path),source_sha256=digest(path),render_status='not_executed')
    if result['status']=='passed':
        powershell=shutil.which('pwsh') or shutil.which('powershell')
        cmd=[powershell,'-NoProfile','-File',str(ROOT/'router/render-powerpoint.ps1'),'-InputPptx',str(path),'-OutDir',str(out/'render')]
        try:
            proc=subprocess.run(cmd,capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=180)
            (out/'render.log').write_text(proc.stdout+proc.stderr,encoding='utf-8')
            images=[out/'render'/f'slide-{i:03}.png' for i in range(1,result['slide_count']+1)]
            if proc.returncode or not all(p.is_file() for p in images): raise RuntimeError('PowerPoint render failed; see render.log')
            contact_sheet(images,out/'contact-sheet.jpg')
            result['render_status']='passed'
            result['render_engine']='Microsoft PowerPoint'
        except Exception as e:
            result['render_status']='failed'; result['errors'].append(str(e)); result['status']='failed'
    result['visual_review']='required'
    write(out/'qa-report.json',result)
    return result

def verify(args):
    report=read(Path(args.qa)/'qa-report.json')
    reviewpath=Path(args.qa)/'visual-review.json'
    review=read(reviewpath) if reviewpath.is_file() else {}
    actual=digest(args.pptx); errors=[]
    if report.get('source_sha256') != actual: errors.append('QA is stale')
    if report.get('status') != 'passed' or report.get('render_status') != 'passed': errors.append('Static/render QA incomplete')
    if review.get('source_sha256') != actual or review.get('status') != 'passed': errors.append('Current visual review missing or needs repair')
    if sorted(review.get('slides_reviewed',[])) != list(range(1,report['slide_count']+1)): errors.append('Visual review must cover all slides')
    result={'status':'failed' if errors else 'passed','errors':errors,'source_sha256':actual,'pptx':str(Path(args.pptx).resolve())}
    write(Path(args.qa)/'final-check.json',result)
    return result

def doctor():
    return {'root':str(ROOT),'python':sys.executable,'node':shutil.which('node'),
        'renderers':{k:(ROOT/'skills-src'/v).is_file() for k,v in RENDERERS.items()},
        'deck_dna':(ROOT/'skills-src/deck-dna/SKILL.md').is_file(),
        'powerpoint':(Path(os.environ.get('ProgramFiles', ''))/'Microsoft Office/root/Office16/POWERPNT.EXE').is_file(),
        'academic_native':'requires configured image provider and upstream TechnicalRoute A/B preflight',
        'discovery_junction':str(Path.home()/'.agents/skills/ppt-router')}

def main():
    p=argparse.ArgumentParser(description=__doc__); sub=p.add_subparsers(dest='cmd',required=True)
    r=sub.add_parser('route'); r.add_argument('--prompt',required=True)
    r.add_argument('--kind',choices=['auto','formal','business','academic','data','simple-course'],default='auto')
    r.add_argument('--slides',type=int,default=5);r.add_argument('--project');r.add_argument('--reference',action='append',default=[])
    r.add_argument('--renderer',choices=list(RENDERERS));r.add_argument('--academic-ready',action='store_true',help='Only after upstream provider/workflow prerequisites are verified and authorized')
    v=sub.add_parser('validate-plan');v.add_argument('path')
    q=sub.add_parser('qa');q.add_argument('--pptx',required=True);q.add_argument('--out',required=True);q.add_argument('--expected-slides',type=int)
    f=sub.add_parser('verify-final');f.add_argument('--pptx',required=True);f.add_argument('--qa',required=True)
    sub.add_parser('doctor')
    args=p.parse_args()
    try:
        result=route(args) if args.cmd=='route' else validate_plan(args.path) if args.cmd=='validate-plan' else qa(args) if args.cmd=='qa' else verify(args) if args.cmd=='verify-final' else doctor()
        print(json.dumps(result,ensure_ascii=False,indent=2))
        return 1 if result.get('status')=='failed' else 0
    except Exception as exc:
        print(json.dumps({'status':'failed','error':str(exc)},ensure_ascii=False));return 1

if __name__=='__main__': raise SystemExit(main())
