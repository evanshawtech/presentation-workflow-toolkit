"""Local read-only deck-dna measurement with grouped text and CJK support.

Upstream's simple split-based counter does not traverse groups or tables and
under-counts Chinese; this helper only extends the measurement at adapter level.
"""
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import statistics
import sys
from pptx import Presentation
from lxml import etree

def walk(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, 'shapes'):
            yield from walk(shape.shapes)

def extract(path):
    prs = Presentation(path)
    colors, fonts, sizes = Counter(), Counter(), Counter()
    counts = []
    for i, slide in enumerate(prs.slides, 1):
        ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main', 'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
        master = slide.slide_layout.slide_master
        theme_rel = next((r for r in master.part.rels.values() if r.reltype.endswith('/theme')), None)
        theme_colors, theme_fonts = {}, {}
        if theme_rel:
            theme = etree.fromstring(theme_rel.target_part.blob)
            scheme = theme.find('a:themeElements/a:clrScheme', ns)
            if scheme is not None:
                for entry in scheme:
                    if len(entry):
                        theme_colors[entry.tag.split('}')[-1]] = entry[0].get('val') if entry[0].tag.endswith('srgbClr') else entry[0].get('lastClr')
            for kind, prefix in [('majorFont', '+mj'), ('minorFont', '+mn')]:
                for script, suffix in [('latin', 'lt'), ('ea', 'ea')]:
                    n = theme.find(f'a:themeElements/a:fontScheme/a:{kind}/a:{script}', ns)
                    if n is not None:
                        theme_fonts[f'{prefix}-{suffix}'] = n.get('typeface', '')
        cmap = master.element.find('p:clrMap', ns)
        color_map = dict(cmap.attrib) if cmap is not None else {}
        texts = []
        for sh in walk(slide.shapes):
            if sh.has_text_frame:
                texts.append(sh.text_frame.text)
            if sh.has_table:
                texts.extend(cell.text for row in sh.table.rows for cell in row.cells)
        text = '\n'.join(texts)
        counts.append({'slide': i, 'whitespace_words': len(text.split()),
                       'cjk_characters': len(re.findall(r'[\u3400-\u9fff]', text)),
                       'nonspace_characters': len(re.sub(r'\s', '', text)),
                       'legacy_top_level_words': sum(len(sh.text_frame.text.split()) for sh in slide.shapes if sh.has_text_frame)})
        for node in slide.element.iter():
            local = node.tag.split('}')[-1]
            if local == 'srgbClr':
                colors['#' + node.get('val', '')] += 1
            elif local == 'schemeClr':
                key = node.get('val', '')
                value = theme_colors.get(color_map.get(key, key))
                if value and not len(node):
                    colors['#' + value] += 1
            elif local in ('latin', 'ea') and node.get('typeface'):
                value = node.get('typeface')
                fonts[theme_fonts.get(value) or value] += 1
            elif local in ('rPr', 'defRPr') and node.get('sz'):
                sizes[str(int(node.get('sz')) / 100)] += 1
    return {'path': str(path.resolve()), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
            'slide_count': len(prs.slides), 'aspect_ratio': round(prs.slide_width / prs.slide_height, 8),
            'slide_counts': counts, 'word_budget': statistics.median(c['whitespace_words'] for c in counts),
            'cjk_character_budget': statistics.median(c['cjk_characters'] for c in counts),
            'character_budget': statistics.median(c['nonspace_characters'] for c in counts),
            'actual_slide_colors': colors.most_common(), 'actual_font_references': fonts.most_common(),
            'font_sizes_pt': sizes.most_common(), 'limitations': ['One reference measured per invocation; combine measurements across approved references before deriving a final budget.',
            'Colors resolve used explicit RGB and unmodified theme references via each slide master; image pixels, color transforms and inherited-only objects require visual confirmation.',
            'Word count includes grouped text and tables; CJK density uses character counts as well.']}

if __name__ == '__main__':
    if len(sys.argv) != 3:
        raise SystemExit('Usage: python extract-design-dna.py reference.pptx output.json')
    result = extract(Path(sys.argv[1]))
    Path(sys.argv[2]).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: result[k] for k in ('slide_count', 'word_budget', 'cjk_character_budget', 'character_budget')}, ensure_ascii=False))
